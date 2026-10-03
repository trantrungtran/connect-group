# -*- coding: utf-8 -*-
"""Build interactive bilingual scripture PowerPoint (sample Mark 7 style)."""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .bible import VerseRef, chunk_by_sections
from .parse_pdf import DiscussionQuestion, StudyContent

# Widescreen 16:9
W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
AMBER = RGBColor(0xC4, 0x8A, 0x3A)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
SOFT = RGBColor(0xE8, 0xEE, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BTN_FILL = RGBColor(0x2F, 0x6F, 0xAD)
BTN_Q = RGBColor(0x8A, 0x4B, 0x2F)
TRIGGER_YELLOW = RGBColor(0xFF, 0xE8, 0x8A)  # pastel yellow
COMMENT_BG = RGBColor(0xFF, 0xFC, 0xF2)  # solid cream — fully opaque panel
COMMENT_TEXT = RGBColor(0x14, 0x14, 0x14)  # near-black, fully opaque text
QUESTION_BG = RGBColor(0xFF, 0xF0, 0xD6)

NSMAP = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


@dataclass
class TriggerPair:
    trigger_spid: int
    target_spid: int


class IdGen:
    def __init__(self, start: int = 2):
        self.n = start

    def next(self) -> int:
        self.n += 1
        return self.n


def _set_run(run, size=14, bold=False, color=NAVY, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                 color=NAVY, font="Calibri", align=PP_ALIGN.LEFT, name=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def _rgb_hex(rgb: RGBColor) -> str:
    return f"{int(rgb[0]):02X}{int(rgb[1]):02X}{int(rgb[2]):02X}"


def _force_opaque_srgb(element, rgb: RGBColor):
    """Ensure this element's solidFill is fully opaque sRGB."""
    for solid in list(element.findall(qn("a:solidFill"))):
        element.remove(solid)
    # Also clear scheme-based fills that can inherit transparency
    for tag in ("a:gradFill", "a:pattFill", "a:noFill"):
        for node in list(element.findall(qn(tag))):
            element.remove(node)
    solid = etree.SubElement(element, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", _rgb_hex(rgb))
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", "100000")


def _add_filled_textbox(slide, left, top, width, height, text, fill, size=13,
                        bold=False, color=NAVY, font="Microsoft YaHei", name=None,
                        opaque=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = AMBER
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    # Comfortable inset so text doesn't sit on the border
    try:
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)

    if opaque:
        spPr = shape._element.spPr
        # Strip soft shadows / glow that can make content look translucent
        for eff in spPr.findall(qn("a:effectLst")):
            spPr.remove(eff)
        _force_opaque_srgb(spPr, fill)
        # Re-apply opaque line color
        ln = spPr.find(qn("a:ln"))
        if ln is None:
            ln = etree.SubElement(spPr, qn("a:ln"))
        _force_opaque_srgb(ln, AMBER)
        # Force every text run fill fully opaque
        for rPr in shape._element.findall(".//" + qn("a:rPr")):
            _force_opaque_srgb(rPr, color)
    return shape


def _add_button(slide, left, top, width, height, label, fill=BTN_FILL, name=None, size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    try:
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    _set_run(run, size=size, bold=True, color=WHITE, font="Calibri")
    return shape


def _add_verse_trigger(
    slide, left, top, width, height, verse_num: int, name=None, alpha_pct=45, font_pt: int = 12
):
    """Pastel-yellow trigger chip with opaque verse-number text drawn on top of the fill."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if name:
        shape.name = name
    # Semi-transparent yellow background
    shape.fill.solid()
    shape.fill.fore_color.rgb = TRIGGER_YELLOW
    spPr = shape._element.spPr
    srgb = spPr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
    if srgb is None:
        srgb = spPr.find(".//" + qn("a:srgbClr"))
    if srgb is not None:
        old = srgb.find(qn("a:alpha"))
        if old is not None:
            srgb.remove(old)
        alpha = etree.SubElement(srgb, qn("a:alpha"))
        alpha.set("val", str(max(0, min(100, int(alpha_pct))) * 1000))
    shape.line.fill.background()

    # Opaque number text sits on the chip (always visually in front of the fill)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    try:
        tf.margin_left = Pt(0)
        tf.margin_right = Pt(0)
        tf.margin_top = Pt(0)
        tf.margin_bottom = Pt(0)
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(verse_num)
    _set_run(run, size=font_pt, bold=True, color=COMMENT_TEXT, font="Calibri")
    for rPr in shape._element.findall(".//" + qn("a:rPr")):
        _force_opaque_srgb(rPr, COMMENT_TEXT)
    # Vertical centering
    try:
        bodyPr = shape._element.txBody.bodyPr
        bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    return shape


def _estimate_wrapped_lines(text: str, col_width_in: float, font_pt: float, cjk: bool) -> int:
    """Rough line count for wrapped verse text (used only for vertical layout)."""
    text = (text or "").strip()
    if not text:
        return 1
    if cjk:
        chars_per_line = max(8, int(col_width_in * 72 / (font_pt * 1.05)))
        # Count CJK as 1, ascii as ~0.55
        units = 0.0
        for ch in text:
            units += 1.0 if "\u4e00" <= ch <= "\u9fff" or ord(ch) > 127 else 0.55
        return max(1, int((units + chars_per_line - 1) // chars_per_line))
    chars_per_line = max(12, int(col_width_in * 72 / (font_pt * 0.52)))
    # Prefer word-aware estimate
    words = text.split()
    lines, cur = 1, 0
    for w in words:
        need = len(w) + (1 if cur else 0)
        if cur + need > chars_per_line and cur:
            lines += 1
            cur = len(w)
        else:
            cur += need
    return max(1, lines)


def _layout_verse_tops(
    nums: List[int],
    en: Dict[int, str],
    zh: Dict[int, str],
    top0: float,
    max_bottom: float,
    en_width_in: float,
    zh_width_in: float,
    font_pt: float = 13,
) -> Dict[int, Tuple[float, float]]:
    """Return verse_num -> (top_inches, block_height_inches) shared across both columns."""
    line_h = font_pt * 1.35 / 72.0  # inches per wrapped line
    gap = 0.12
    # First pass: natural heights
    heights = {}
    for n in nums:
        en_lines = _estimate_wrapped_lines(f"{n}  {en.get(n, '')}", en_width_in - 0.05, font_pt, cjk=False)
        zh_lines = _estimate_wrapped_lines(f"{n}  {zh.get(n, '')}", zh_width_in - 0.05, font_pt, cjk=True)
        heights[n] = max(en_lines, zh_lines) * line_h + gap

    natural = sum(heights.values())
    usable = max(0.5, max_bottom - top0)
    scale = min(1.0, usable / natural) if natural else 1.0

    tops: Dict[int, Tuple[float, float]] = {}
    cursor = top0
    for n in nums:
        h = heights[n] * scale
        tops[n] = (cursor, h)
        cursor += h
    return tops


def _add_verse_column(
    slide,
    nums: List[int],
    texts: Dict[int, str],
    layout: Dict[int, Tuple[float, float]],
    left_in: float,
    width_in: float,
    font: str,
    name_prefix: str,
    size: int = 13,
    number_gutter_in: float = 0.0,
    numbered_verses: Optional[set] = None,
):
    """Place one textbox per verse so number Y is known exactly.

    If a verse is in numbered_verses, its number is drawn on the trigger chip instead,
    so the body text is indented by number_gutter_in and omits the leading number.
    """
    numbered_verses = numbered_verses or set()
    for n in nums:
        top, h = layout[n]
        body = texts.get(n, "").strip()
        if n in numbered_verses:
            text = body
            left = left_in + number_gutter_in
            width = max(0.5, width_in - number_gutter_in)
        else:
            text = f"{n}  {body}" if body else str(n)
            left = left_in
            width = width_in
        if not text:
            continue
        _add_textbox(
            slide,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(max(0.22, h - 0.02)),
            text,
            size=size,
            color=NAVY,
            font=font,
            name=f"{name_prefix} {n}",
        )


def _hotspot_width_for_verse(n: int) -> float:
    """Width in inches to cover the verse number digits (+ small pad)."""
    digits = len(str(n))
    return 0.18 + 0.12 * digits


def _set_visibility(parent, idgen: IdGen, spid: int, visible: bool, delay: str = "0"):
    """Append a style.visibility set animation for one shape."""
    P = NSMAP["p"]
    set_el = etree.SubElement(parent, f"{{{P}}}set")
    bh = etree.SubElement(set_el, f"{{{P}}}cBhvr")
    ctn = etree.SubElement(bh, f"{{{P}}}cTn", id=str(idgen.next()), dur="1", fill="hold")
    stc = etree.SubElement(ctn, f"{{{P}}}stCondLst")
    etree.SubElement(stc, f"{{{P}}}cond", delay=delay)
    t = etree.SubElement(bh, f"{{{P}}}tgtEl")
    etree.SubElement(t, f"{{{P}}}spTgt", spid=str(spid))
    attrs = etree.SubElement(bh, f"{{{P}}}attrNameLst")
    etree.SubElement(attrs, f"{{{P}}}attrName").text = "style.visibility"
    to = etree.SubElement(set_el, f"{{{P}}}to")
    etree.SubElement(to, f"{{{P}}}strVal", val="visible" if visible else "hidden")


def _make_fly_seq(
    idgen: IdGen,
    trigger_spid: int,
    target_spid: int,
    entrance: bool,
) -> etree._Element:
    """Fly in/out target on click. Triggers stay visible; z-order keeps overlays on top."""
    P = NSMAP["p"]
    seq = etree.Element(f"{{{P}}}seq", concurrent="1", nextAc="seek")
    cTn = etree.SubElement(
        seq,
        f"{{{P}}}cTn",
        id=str(idgen.next()),
        restart="whenNotActive",
        fill="hold",
        evtFilter="cancelBubble",
        nodeType="interactiveSeq",
    )
    st = etree.SubElement(cTn, f"{{{P}}}stCondLst")
    cond = etree.SubElement(st, f"{{{P}}}cond", evt="onClick", delay="0")
    tgt = etree.SubElement(cond, f"{{{P}}}tgtEl")
    etree.SubElement(tgt, f"{{{P}}}spTgt", spid=str(trigger_spid))

    endSync = etree.SubElement(cTn, f"{{{P}}}endSync", evt="end", delay="0")
    etree.SubElement(endSync, f"{{{P}}}rtn", val="all")

    child = etree.SubElement(cTn, f"{{{P}}}childTnLst")
    par1 = etree.SubElement(child, f"{{{P}}}par")
    c1 = etree.SubElement(par1, f"{{{P}}}cTn", id=str(idgen.next()), fill="hold")
    st1 = etree.SubElement(c1, f"{{{P}}}stCondLst")
    etree.SubElement(st1, f"{{{P}}}cond", delay="0")
    ch1 = etree.SubElement(c1, f"{{{P}}}childTnLst")

    par2 = etree.SubElement(ch1, f"{{{P}}}par")
    c2 = etree.SubElement(par2, f"{{{P}}}cTn", id=str(idgen.next()), fill="hold")
    st2 = etree.SubElement(c2, f"{{{P}}}stCondLst")
    etree.SubElement(st2, f"{{{P}}}cond", delay="0")
    ch2 = etree.SubElement(c2, f"{{{P}}}childTnLst")

    par3 = etree.SubElement(ch2, f"{{{P}}}par")
    preset_class = "entr" if entrance else "exit"
    c3 = etree.SubElement(
        par3,
        f"{{{P}}}cTn",
        id=str(idgen.next()),
        presetID="2",
        presetClass=preset_class,
        presetSubtype="4",
        fill="hold",
        grpId="0" if entrance else "2",
        nodeType="clickEffect",
    )
    st3 = etree.SubElement(c3, f"{{{P}}}stCondLst")
    etree.SubElement(st3, f"{{{P}}}cond", delay="0")
    ch3 = etree.SubElement(c3, f"{{{P}}}childTnLst")

    def add_anim_y(start_val: str, end_val: str, dur: str = "500"):
        anim = etree.SubElement(ch3, f"{{{P}}}anim", calcmode="lin", valueType="num")
        bh = etree.SubElement(anim, f"{{{P}}}cBhvr", additive="base")
        etree.SubElement(bh, f"{{{P}}}cTn", id=str(idgen.next()), dur=dur, fill="hold")
        t = etree.SubElement(bh, f"{{{P}}}tgtEl")
        etree.SubElement(t, f"{{{P}}}spTgt", spid=str(target_spid))
        attrs = etree.SubElement(bh, f"{{{P}}}attrNameLst")
        etree.SubElement(attrs, f"{{{P}}}attrName").text = "ppt_y"
        tavs = etree.SubElement(anim, f"{{{P}}}tavLst")
        for tm, val in (("0", start_val), ("100000", end_val)):
            tav = etree.SubElement(tavs, f"{{{P}}}tav", tm=tm)
            v = etree.SubElement(tav, f"{{{P}}}val")
            etree.SubElement(v, f"{{{P}}}strVal", val=val)

    if entrance:
        _set_visibility(ch3, idgen, target_spid, True, delay="0")
        add_anim_y("1+#ppt_h/2", "#ppt_y")
    else:
        add_anim_y("ppt_y", "1+ppt_h/2")
        _set_visibility(ch3, idgen, target_spid, False, delay="499")

    nxt = etree.SubElement(seq, f"{{{P}}}nextCondLst")
    ncond = etree.SubElement(nxt, f"{{{P}}}cond", evt="onClick", delay="0")
    nt = etree.SubElement(ncond, f"{{{P}}}tgtEl")
    etree.SubElement(nt, f"{{{P}}}spTgt", spid=str(trigger_spid))
    return seq


# (trigger_spid, target_spid)
AnimPair = Tuple[int, int]


def _build_timing_xml(pairs: List[AnimPair]) -> etree._Element:
    P = NSMAP["p"]
    timing = etree.Element(f"{{{P}}}timing")
    tnLst = etree.SubElement(timing, f"{{{P}}}tnLst")
    par = etree.SubElement(tnLst, f"{{{P}}}par")
    root = etree.SubElement(
        par, f"{{{P}}}cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot"
    )
    child = etree.SubElement(root, f"{{{P}}}childTnLst")
    idgen = IdGen(1)
    for trig, target in pairs:
        child.append(_make_fly_seq(idgen, trig, target, entrance=True))
        child.append(_make_fly_seq(idgen, target, target, entrance=False))
    return timing


def _attach_timing(slide, pairs: List[AnimPair]):
    if not pairs:
        return
    timing = _build_timing_xml(pairs)
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    sld.append(timing)


def _shape_cNvPr_name(el) -> str:
    for node in el.iter():
        if node.tag == qn("p:cNvPr") or node.tag.endswith("}cNvPr"):
            return node.get("name") or ""
    return ""


def _finalize_slide_zorder(slide):
    """Force shape tree paint order: text → controls → commentaries → question popups.

    PowerPoint Selection pane lists front-to-back (top = front), which is the reverse
    of spTree sibling order (last sibling = front).
    """
    spTree = slide.shapes._spTree
    headers = []
    shapes = []
    for child in list(spTree):
        tag = child.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            headers.append(child)
        else:
            shapes.append(child)

    def layer(el) -> int:
        name = _shape_cNvPr_name(el)
        if name == "Title" or name.startswith("English ") or name.startswith("Chinese "):
            return 0
        if name.startswith("Question:"):
            return 3
        if "Comment" in name:
            return 2
        # Verse triggers and discussion buttons stay under overlays
        return 1

    shapes.sort(key=lambda el: (layer(el), _shape_cNvPr_name(el)))
    for child in list(spTree):
        spTree.remove(child)
    for h in headers:
        spTree.append(h)
    for s in shapes:
        spTree.append(s)


def _assign_questions_to_chunks(
    questions: List[DiscussionQuestion], n_chunks: int
) -> List[List[DiscussionQuestion]]:
    """Place each question on its author section slide (section_index).

    Falls back to even spread only for questions with no section assignment.
    """
    out: List[List[DiscussionQuestion]] = [[] for _ in range(max(n_chunks, 1))]
    if not questions:
        return out

    unassigned: List[DiscussionQuestion] = []
    for q in questions:
        idx = getattr(q, "section_index", -1)
        if 0 <= idx < n_chunks:
            out[idx].append(q)
        else:
            unassigned.append(q)

    if unassigned and n_chunks > 0:
        # Prefer empty section slides first, then append evenly
        for i, q in enumerate(unassigned):
            empty = next((j for j in range(n_chunks) if not out[j]), None)
            if empty is not None:
                out[empty].append(q)
            else:
                out[min(i % n_chunks, n_chunks - 1)].append(q)

    # Keep question number order within each slide
    for bucket in out:
        bucket.sort(key=lambda q: q.number)
    return out


def build_presentation(
    study: StudyContent,
    niv: Dict[int, str],
    cuv: Dict[int, str],
    out_path: str,
) -> str:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    ref = study.ref
    assert ref is not None, "Missing scripture reference"
    title_en = getattr(study, "title_en", "") or study.title_zh
    verse_nums = [n for n in ref.verse_numbers() if niv.get(n) or cuv.get(n) or n in study.commentaries_zh]
    if not verse_nums:
        verse_nums = ref.verse_numbers()
    # Prefer author section breaks (e.g. 8:1-10 / 8:11-13 / 8:14-21)
    section_chunks = chunk_by_sections(
        verse_nums, getattr(study, "sections", []) or [], max_per_slide=10
    )
    q_chunks = _assign_questions_to_chunks(study.questions, len(section_chunks))

    for ci, (nums, sec_title_zh, sec_title_en) in enumerate(section_chunks):
        slide = prs.slides.add_slide(blank)
        pairs: List[AnimPair] = []

        # Title: verse range + author section title when available
        if nums[0] == nums[-1]:
            range_label = f"{ref.book_en} {ref.chapter}:{nums[0]}"
        else:
            range_label = f"{ref.book_en} {ref.chapter}:{nums[0]}–{nums[-1]}"
        sec_label = sec_title_zh or study.title_zh
        if sec_title_en and sec_title_zh:
            header = f"{range_label}   ·   {sec_title_en} / {sec_title_zh}"
        else:
            header = f"{range_label}   ·   {sec_label}"
        # Keep header readable on one line
        if len(header) > 70:
            header = f"{range_label}   ·   {sec_title_zh or study.title_zh}"
        _add_textbox(
            slide,
            Inches(0.5),
            Inches(0.08),
            Inches(12.3),
            Inches(0.45),
            header,
            size=20 if len(header) > 55 else 22,
            bold=True,
            color=NAVY,
            font="Georgia",
            name="Title",
        )

        # Per-verse scripture layout; numbers live on trigger chips (opaque text on yellow)
        en_left, zh_left = 0.35, 6.70
        en_width, zh_width = 6.10, 6.20
        text_top, text_bottom = 0.55, 6.70
        font_pt = 13
        numbered = {
            n
            for n in nums
            if study.commentaries_zh.get(n) or study.commentaries_en.get(n)
        }
        # Gutter wide enough for 1–2 digit chips + small gap before verse text
        number_gutter = 0.42
        layout = _layout_verse_tops(
            nums, niv, cuv, text_top, text_bottom, en_width - number_gutter, zh_width - number_gutter, font_pt=font_pt
        )
        _add_verse_column(
            slide,
            nums,
            niv,
            layout,
            en_left,
            en_width,
            "Calibri",
            "English",
            size=font_pt,
            number_gutter_in=number_gutter,
            numbered_verses=numbered,
        )
        _add_verse_column(
            slide,
            nums,
            cuv,
            layout,
            zh_left,
            zh_width,
            "Microsoft YaHei",
            "Chinese",
            size=font_pt,
            number_gutter_in=number_gutter,
            numbered_verses=numbered,
        )

        # Controls first, overlays last; z-order keeps overlays on top (triggers stay visible)
        hotspot_h = 0.24
        y_nudge = 0.01
        pending_comments = []

        for n in nums:
            zh_c = study.commentaries_zh.get(n, "")
            en_c = study.commentaries_en.get(n, "")
            if not zh_c and not en_c:
                continue

            body = zh_c
            if en_c:
                body = f"{en_c}\n\n{zh_c}" if zh_c else en_c
            h_in = min(2.4, max(0.7, 0.35 + len(body) / 120))
            verse_top, _ = layout[n]
            top_in = min(max(verse_top, text_top), 7.0 - h_in - 0.55)
            hotspot_w = _hotspot_width_for_verse(n)
            hotspot_top = verse_top + y_nudge

            btn_en = _add_verse_trigger(
                slide,
                Inches(en_left),
                Inches(hotspot_top),
                Inches(hotspot_w),
                Inches(hotspot_h),
                verse_num=n,
                name=f"{n} in English",
                font_pt=12,
            )
            btn_zh = _add_verse_trigger(
                slide,
                Inches(zh_left),
                Inches(hotspot_top),
                Inches(hotspot_w),
                Inches(hotspot_h),
                verse_num=n,
                name=f"{n} in",
                font_pt=12,
            )
            pending_comments.append(
                {
                    "n": n,
                    "body": body,
                    "top_in": top_in,
                    "h_in": h_in,
                    "btn_en": btn_en.shape_id,
                    "btn_zh": btn_zh.shape_id,
                }
            )

        qs = q_chunks[ci] if ci < len(q_chunks) else []
        pending_questions = []
        left = 0.18
        for q in qs:
            label = f"{q.number}) {q.short_label or 'Q'}?"
            w = max(1.1, min(2.2, 0.55 + 0.11 * len(label)))
            btn = _add_button(
                slide,
                Inches(left),
                Inches(6.85),
                Inches(w),
                Inches(0.42),
                label,
                fill=BTN_Q,
                name=f"Q {q.short_label} in",
                size=11,
            )
            left += w + 0.12
            qtext = q.english
            if q.chinese:
                qtext = f"{q.english}\n\n{q.chinese}" if q.english else q.chinese
            pending_questions.append(
                {
                    "btn_id": btn.shape_id,
                    "label": q.short_label or f"Q{q.number}",
                    "text": f"{q.short_label}: {qtext}",
                }
            )

        for item in pending_comments:
            n = item["n"]
            comment_right = _add_filled_textbox(
                slide,
                Inches(zh_left - 0.05),
                Inches(item["top_in"]),
                Inches(6.25),
                Inches(item["h_in"]),
                item["body"],
                fill=COMMENT_BG,
                size=12,
                color=COMMENT_TEXT,
                font="Microsoft YaHei",
                name=f"{n} Comment Right",
                opaque=True,
            )
            pairs.append((item["btn_en"], comment_right.shape_id))

            comment_left = _add_filled_textbox(
                slide,
                Inches(0.20),
                Inches(item["top_in"]),
                Inches(6.25),
                Inches(item["h_in"]),
                item["body"],
                fill=COMMENT_BG,
                size=12,
                color=COMMENT_TEXT,
                font="Microsoft YaHei",
                name=f"{n} Comment Left",
                opaque=True,
            )
            pairs.append((item["btn_zh"], comment_left.shape_id))

        for item in pending_questions:
            qbox = _add_filled_textbox(
                slide,
                Inches(3.2),
                Inches(5.7),
                Inches(9.7),
                Inches(1.35),
                item["text"],
                fill=QUESTION_BG,
                size=13,
                color=COMMENT_TEXT,
                font="Microsoft YaHei",
                name=f"Question: {item['label']}",
                opaque=True,
            )
            pairs.append((item["btn_id"], qbox.shape_id))

        _finalize_slide_zorder(slide)
        _attach_timing(slide, pairs)

    # Optional closing title meta slide? User asked to focus on bible text — skip extras.
    prs.save(out_path)
    return out_path
