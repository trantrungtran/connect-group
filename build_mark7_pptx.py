# -*- coding: utf-8 -*-
"""Build ROLCC Connect Group PowerPoint from Mark 7 celldoc."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand palette — deep navy + warm amber (readable on projector)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
NAVY_DARK = RGBColor(0x12, 0x1C, 0x33)
AMBER = RGBColor(0xC4, 0x8A, 0x3A)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x3A, 0x45, 0x5C)
SOFT = RGBColor(0xE8, 0xEE, 0xF5)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size=18, bold=False, color=NAVY, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                color=NAVY, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_paras(slide, left, top, width, height, lines, size=16, color=SLATE,
              bold_first=False, spacing=6, font="Calibri"):
    """lines: list of str, or list of (text, bold, size_override)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, is_bold, sz = item[0], item[1], item[2] if len(item) > 2 else size
        else:
            text, is_bold, sz = item, (bold_first and i == 0), size
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = text
        set_run(run, size=sz, bold=is_bold, color=color, font=font)
    return box


def header_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, W, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), W, Pt(4), AMBER)
    add_textbox(slide, Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.5),
                title, size=26, bold=True, color=WHITE, font="Georgia")
    if subtitle:
        add_textbox(slide, Inches(0.5), Inches(0.65), Inches(12.3), Inches(0.35),
                    subtitle, size=12, color=Rgb, font="Calibri")


def footer(slide, page, total):
    add_textbox(slide, Inches(0.5), Inches(7.1), Inches(8), Inches(0.3),
                "ROLCC Connect Group  ·  Mark Series 14", size=10, color=Rgb)
    add_textbox(slide, Inches(10.5), Inches(7.1), Inches(2.3), Inches(0.3),
                f"{page} / {total}", size=10, color=SLATE, align=PP_ALIGN.RIGHT)


def two_col_scripture(slide, en_lines, cn_lines):
    # left panel
    add_rect(slide, Inches(0.4), Inches(1.3), Inches(6.1), Inches(5.5), SOFT)
    add_textbox(slide, Inches(0.6), Inches(1.45), Inches(5.7), Inches(0.35),
                "NIV (English)", size=12, bold=True, color=AMBER)
    add_paras(slide, Inches(0.6), Inches(1.85), Inches(5.7), Inches(4.7),
              en_lines, size=13, color=NAVY, spacing=8)
    # right panel
    add_rect(slide, Inches(6.8), Inches(1.3), Inches(6.1), Inches(5.5), CREAM)
    add_textbox(slide, Inches(7.0), Inches(1.45), Inches(5.7), Inches(0.35),
                "和合本 (Chinese)", size=12, bold=True, color=AMBER)
    add_paras(slide, Inches(7.0), Inches(1.85), Inches(5.7), Inches(4.7),
              cn_lines, size=13, color=NAVY, spacing=8, font="Microsoft YaHei")


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    slides_meta = []  # (builder fn) applied after we know total — we'll build then stamp footers

    # ---- 1 Title ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY_DARK)
    add_rect(s, 0, Inches(5.9), W, Inches(1.6), NAVY)
    add_rect(s, Inches(0.7), Inches(2.0), Inches(1.2), Pt(5), AMBER)
    add_textbox(s, Inches(0.7), Inches(2.2), Inches(12), Inches(0.45),
                "ROLCC Connect Group", size=16, color=AMBER, font="Calibri")
    add_textbox(s, Inches(0.7), Inches(2.7), Inches(12), Inches(1.2),
                "Mark Series 14", size=44, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(3.9), Inches(12), Inches(0.9),
                "Jesus Casts Out a Demon and\nHeals a Deaf and Mute Man", size=24, color=CREAM, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(6.2), Inches(12), Inches(0.4),
                "Mark 7:24–37  ·  September 18–19, 2026", size=16, color=SOFT)
    add_textbox(s, Inches(0.7), Inches(6.65), Inches(12), Inches(0.35),
                "馬可福音系列 14 — 耶穌趕鬼並醫治耳聾舌結的人", size=14, color=AMBER, font="Microsoft YaHei")

    # ---- 2 Opening ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Opening Discussion", "討論開場")
    add_rect(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(2.2), SOFT)
    add_textbox(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.4),
                "1. Last Sunday’s message", size=14, bold=True, color=AMBER)
    add_textbox(s, Inches(0.8), Inches(2.15), Inches(11.7), Inches(1.2),
                "What was the central thought of last Sunday’s message?\nWhich sentence left the deepest impression on you?\n\n上主日信息的中心思想是什麼？哪一句話給你最深的印象？",
                size=18, color=NAVY, font="Microsoft YaHei")
    add_rect(s, Inches(0.5), Inches(4.0), Inches(12.3), Inches(2.5), CREAM)
    add_textbox(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.4),
                "2. Response & action", size=14, bold=True, color=AMBER)
    add_textbox(s, Inches(0.8), Inches(4.65), Inches(11.7), Inches(1.5),
                "After hearing the message, how do you respond?\nWhat practical action have you taken or will take?\n\n聽了信息之後，你如何回應？你已經或將要採取什麼實際行動？",
                size=18, color=NAVY, font="Microsoft YaHei")

    # ---- 3 Scripture 24-30 ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Scripture — Mark 7:24–30", "經文：馬可福音 7:24–30")
    en = [
        "24 Jesus left that place and went to the vicinity of Tyre. He entered a house and did not want anyone to know it; yet he could not keep his presence secret.",
        "25 In fact, as soon as she heard about him, a woman whose little daughter was possessed by an impure spirit came and fell at his feet.",
        "26 The woman was a Greek, born in Syrian Phoenicia. She begged Jesus to drive the demon out of her daughter.",
        "27 “First let the children eat all they want,” he told her, “for it is not right to take the children’s bread and toss it to the dogs.”",
        "28 “Lord,” she replied, “even the dogs under the table eat the children’s crumbs.”",
        "29 Then he told her, “For such a reply, you may go; the demon has left your daughter.”",
        "30 She went home and found her child lying on the bed, and the demon gone.",
    ]
    cn = [
        "24 耶穌從那裡起身，往推羅、西頓的境內去，進了一家，不願意人知道，卻隱藏不住。",
        "25 當下，有一個婦人，他的小女兒被污鬼附著，聽見耶穌的事，就來俯伏在他腳前。",
        "26 這婦人是希利尼人，屬敘利腓尼基族。他求耶穌趕出那鬼離開他的女兒。",
        "27 耶穌對他說：讓兒女們先吃飽，不好拿兒女的餅丟給狗吃。",
        "28 婦人回答說：主阿，不錯；但是狗在桌子底下也吃孩子們的碎渣兒。",
        "29 耶穌對他說：因這句話，你回去罷；鬼已經離開你的女兒了。",
        "30 他就回家去，見小孩子躺在床上，鬼已經出去了。",
    ]
    two_col_scripture(s, en, cn)

    # ---- 4 Section I ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY)
    add_textbox(s, Inches(0.7), Inches(2.5), Inches(12), Inches(0.4),
                "Part I", size=16, color=AMBER)
    add_textbox(s, Inches(0.7), Inches(3.0), Inches(12), Inches(1.2),
                "The Gentile Woman’s Faith\n外邦婦人的信心", size=36, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(4.6), Inches(12), Inches(0.5),
                "Mark 7:24–30", size=18, color=SOFT)

    # ---- 5 A ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "A. The Gentile woman begs Jesus", "外邦婦人懇求耶穌（7:24–26）")
    points = [
        ("Tyre & Sidon — Gentile territory", False, 18),
        ("Isaiah 23 & Joel 3 record God’s curse on these places — yet Jesus goes there quietly.", False, 15),
        ("A Syrophoenician (Greek) woman’s little daughter was possessed by an unclean spirit.", False, 15),
        ("Hearing of Jesus, she came in faith, fell at His feet, and begged Him to drive the demon out.", False, 15),
        ("", False, 10),
        ("推羅、西頓是外邦之地；以賽亞書23、約珥書3記載神的咒詛，耶穌仍悄悄前去。", False, 15),
        ("敘利腓尼基婦人的小女兒被污鬼附著；她憑信心來到耶穌腳前懇求趕鬼。", False, 15),
    ]
    add_paras(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5.2), points, size=16, color=NAVY, spacing=10, font="Microsoft YaHei")

    # ---- 6 B ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "B. The Gentile woman’s humility", "外邦婦人的謙卑（7:27–28）")
    add_rect(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(1.6), SOFT)
    add_textbox(s, Inches(0.8), Inches(1.55), Inches(11.7), Inches(1.3),
                "“First let the children eat all they want… it is not right to take the children’s bread and toss it to the dogs.”\n「讓兒女們先吃飽，不好拿兒女的餅丟給狗吃。」",
                size=16, color=NAVY, font="Microsoft YaHei")
    points = [
        "Jews believed God’s grace was first for His chosen people, then to the Gentiles.",
        "She did not take offense. She humbly seized the opportunity:",
        "“Lord… even the dogs under the table eat the children’s crumbs.”",
        "She believed crumbs of grace were enough — and fitting — for her daughter.",
        "",
        "她沒有被冒犯，反而謙卑抓住機會：連狗也能吃到孩子們的碎渣兒。",
        "她相信外邦人也能領受耶穌醫治與趕鬼的恩典。",
    ]
    add_paras(s, Inches(0.7), Inches(3.3), Inches(12), Inches(3.5), points, size=16, color=NAVY, spacing=8, font="Microsoft YaHei")

    # ---- 7 C ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "C. The daughter is healed", "外邦婦人的女兒得醫治（7:29–30）")
    add_rect(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(4.8), CREAM)
    points = [
        ("Jesus saw both faith and humble, persistent pleading.", True, 20),
        ("Without going to her house, He drove the demon out.", False, 18),
        ("“For such a reply, you may go; the demon has left your daughter.”", False, 18),
        ("She went home and found her child on the bed — the demon gone.", False, 18),
        ("", False, 12),
        ("耶穌看見她有信心又謙卑地緊緊懇求；雖未到她家，鬼已離開。", False, 17),
        ("「因這句話，你回去罷；鬼已經離開你的女兒了。」", False, 17),
    ]
    add_paras(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(4.2), points, size=18, color=NAVY, spacing=12, font="Microsoft YaHei")

    # ---- 8 Q2 Race ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Discussion 2 — Race / Territory", "討論二：族群／地域")
    add_rect(s, Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.5), SOFT)
    add_textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(3.5),
                "Why did Jesus leave the Jewish region and go into Gentile territory?\nIf you had been one of the disciples following Him, what would you have thought?\n\n為何耶穌離開猶太人地區，到外邦人的境內去？\n你若是跟隨祂的門徒之一，會有什麼想法？\n\n(Observation / Application · 觀察／應用)",
                size=20, color=NAVY, font="Microsoft YaHei")

    # ---- 9 Q3 Faith ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Discussion 3 — Faith", "討論三：信心")
    add_rect(s, Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.5), CREAM)
    add_textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(3.2),
                "From this Gentile woman, how do we see that she was full of faith?\n\n從這外邦婦人的身上，如何看到她滿有信心？\n\n(Observation · 觀察)",
                size=22, color=NAVY, font="Microsoft YaHei")

    # ---- 10 Themes ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Reflection Themes", "反思主題")
    cards = [
        (0.5, "Fear of the Lord\n敬畏主",
         "Children can be possessed.\nFear of danger without the Shepherd’s protection.\nFear of missing out (FOMO).\n\n孩子也可能被鬼附！\n沒有牧人保護的懼怕；害怕錯過。"),
        (4.55, "Possession\n附身",
         "Suggested a lie → believed a lie → living the lie → preaching the lie → automatic acts → yielding to the accuser’s voice / control.\n\n聽見謊言、相信、活出、傳揚、習慣而行、向說謊者降服。"),
        (8.6, "Humility\n謙卑",
         "Compare the Gentile woman with the Jews.\nWith what attitude should we come to receive God’s grace?\n\n我們當用什麼態度來領受神的恩典？"),
    ]
    for left, title, body in cards:
        add_rect(s, Inches(left), Inches(1.4), Inches(3.85), Inches(5.2), SOFT)
        add_rect(s, Inches(left), Inches(1.4), Inches(3.85), Inches(1.15), NAVY)
        add_textbox(s, Inches(left + 0.2), Inches(1.55), Inches(3.45), Inches(0.9),
                    title, size=16, bold=True, color=WHITE, font="Microsoft YaHei")
        add_textbox(s, Inches(left + 0.2), Inches(2.75), Inches(3.45), Inches(3.5),
                    body, size=13, color=NAVY, font="Microsoft YaHei")

    # ---- 11 Scripture 31-37 ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Scripture — Mark 7:31–37", "經文：馬可福音 7:31–37")
    en2 = [
        "31 Then Jesus left the vicinity of Tyre and went through Sidon, down to the Sea of Galilee and into the region of the Decapolis.",
        "32 There some people brought to him a man who was deaf and could hardly talk, and they begged Jesus to place his hand on him.",
        "33 After he took him aside, away from the crowd, Jesus put his fingers into the man’s ears. Then he spit and touched the man’s tongue.",
        "34 He looked up to heaven and with a deep sigh said to him, “Ephphatha!” (which means “Be opened!”).",
        "35 At this, the man’s ears were opened, his tongue was loosened and he began to speak plainly.",
        "36 Jesus commanded them not to tell anyone. But the more he did so, the more they kept talking about it.",
        "37 People were overwhelmed with amazement. “He has done everything well,” they said. “He even makes the deaf hear and the mute speak.”",
    ]
    cn2 = [
        "31 耶穌又離了推羅的境界，經過西頓，就從低加波利境內來到加利利海。",
        "32 有人帶著一個耳聾舌結的人來見耶穌，求他按手在他身上。",
        "33 耶穌領他離開眾人，到一邊去，就用指頭探他的耳朵，吐唾沫抹他的舌頭，",
        "34 望天歎息，對他說：以法大！就是說：開了罷！",
        "35 他的耳朵就開了，舌結也解了，說話也清楚了。",
        "36 耶穌囑咐他們不要告訴人；但他越發囑咐，他們越發傳揚開了。",
        "37 眾人分外希奇，說：他所作的事都好，他連聾子也叫他們聽見，啞巴也叫他們說話。",
    ]
    two_col_scripture(s, en2, cn2)

    # ---- 12 Section II ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY)
    add_textbox(s, Inches(0.7), Inches(2.5), Inches(12), Inches(0.4),
                "Part II", size=16, color=AMBER)
    add_textbox(s, Inches(0.7), Inches(3.0), Inches(12), Inches(1.2),
                "Jesus Heals a Deaf and Mute Man\n耶穌醫治耳聾舌結的人", size=34, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(4.6), Inches(12), Inches(0.5),
                "Mark 7:31–37", size=18, color=SOFT)

    # ---- 13 A ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "A. A deaf and mute man is brought", "有人帶來耳聾舌結的人（7:31–32）")
    points = [
        "Jesus left Tyre, went through Sidon and the Decapolis (Ten Cities — Greco-Roman), to the Sea of Galilee.",
        "People brought a man who was deaf and had a speech impediment.",
        "They specifically asked Jesus to lay His hands on him — perhaps assuming that was always His method.",
        "",
        "耶穌經低加波利來到加利利海（猶太人居住地區）。",
        "有人帶耳聾舌結的人來，特別求耶穌按手醫治——可能已有先入為主的觀念。",
    ]
    add_paras(s, Inches(0.7), Inches(1.6), Inches(12), Inches(5), points, size=17, color=NAVY, spacing=12, font="Microsoft YaHei")

    # ---- 14 B ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "B. Jesus’ unusual method", "耶穌獨特的醫治方法（7:33–35）")
    points = [
        ("He did not follow their request.", True, 18),
        ("Took the man aside → fingers in ears → spit → touched tongue.", False, 17),
        ("Looked up to heaven, sighed, said “Ephphatha!” — “Be opened!”", False, 17),
        ("At once: ears opened, tongue loosened, he spoke plainly.", False, 17),
        ("", False, 10),
        ("The sigh may reflect Jewish lack of faith / self-assurance — knowing Jesus can heal, yet not entrusting Lordship to Him — unlike the Gentile woman’s faith and humility.", False, 15),
        ("", False, 8),
        ("耶穌沒按他們的要求，用獨特方法；望天歎息說「以法大——開了罷！」", False, 15),
        ("歎息可能因猶太人缺乏信心或自以為是——知道耶穌能醫治，卻未將主權交給祂。", False, 15),
    ]
    add_paras(s, Inches(0.7), Inches(1.45), Inches(12), Inches(5.3), points, size=16, color=NAVY, spacing=8, font="Microsoft YaHei")

    # ---- 15 C ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "C. The crowd spreads His fame", "眾人傳揚耶穌（7:36–37）")
    points = [
        "Jesus charged them not to tell anyone — yet the more He charged, the more they proclaimed.",
        "The crowd was astonished: “He has done everything well… He even makes the deaf hear and the mute speak.”",
        "Fulfillment of Isaiah 35:5 — “Then will the eyes of the blind be opened and the ears of the deaf unstopped…”",
        "",
        "耶穌越囑咐不要告訴人，他們越發傳揚。",
        "眾人希奇：「他所做的事都好……」——應驗以賽亞書 35:5 的預言。",
    ]
    add_paras(s, Inches(0.7), Inches(1.6), Inches(12), Inches(5), points, size=17, color=NAVY, spacing=12, font="Microsoft YaHei")

    # ---- 16 Q4 ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Discussion 4 — Attitude", "討論四：態度")
    add_rect(s, Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.5), SOFT)
    add_textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(3.5),
                "Compare the attitude of the Gentile woman with that of the Jews.\nWith what attitude should we come to receive God’s grace?\n\n請對比外邦婦人和猶太人的態度。\n我們當用什麼態度來領受神的恩典？\n\n(Observation / Application · 觀察／應用)",
                size=20, color=NAVY, font="Microsoft YaHei")

    # ---- 17 Q5 ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Discussion 5 — Method & Healing", "討論五：方法與醫治")
    add_rect(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(2.5), CREAM)
    add_textbox(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(2.1),
                "Method: Why did Jesus use an unusual method to heal the deaf and mute man?\nPlease share a testimony of receiving God’s healing.\n\n耶穌為何用獨特的方法醫治耳聾舌結的人？請分享得神醫治的見證。",
                size=17, color=NAVY, font="Microsoft YaHei")
    add_rect(s, Inches(0.5), Inches(4.3), Inches(12.3), Inches(2.2), SOFT)
    add_textbox(s, Inches(0.9), Inches(4.55), Inches(11.5), Inches(1.7),
                "Healing: How can we learn to heal like Jesus?\n我們怎能學像耶穌那樣醫治？",
                size=18, color=NAVY, font="Microsoft YaHei")

    # ---- 18 Takeaway ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Key Takeaway", "核心信息")
    points = [
        ("God’s grace is not only for His chosen people — it is also for the Gentiles.", True, 18),
        ("People must still come before God to receive His grace.", False, 17),
        ("Come with faith and humility — ask Him for the grace He has prepared.", False, 17),
        ("Jesus meets needs — sometimes in ways we never expected.", False, 17),
        ("", False, 10),
        ("神的恩典不僅給選民，也給外邦人；人必須來到神面前才能領受。", False, 16),
        ("當存信心與謙卑來到祂面前，求祂賜下已為我們預備的恩典。", False, 16),
        ("阿們。", True, 18),
    ]
    add_paras(s, Inches(0.7), Inches(1.6), Inches(12), Inches(5), points, size=17, color=NAVY, spacing=10, font="Microsoft YaHei")

    # ---- 19 Prayer ----
    s = prs.slides.add_slide(blank)
    header_bar(s, "Prayer", "禱告")
    prayer_en = (
        "Dear Abba Father, we thank You and praise You! Thank You for teaching us that whoever comes to Jesus with a need and asks will receive, "
        "and that He will accomplish everything — even more than we ask or imagine. Therefore we should learn from the Gentile woman: come before You "
        "with a heart that trusts You, come humbly, ask You, and hand the lordship over to You. You will do wonderful things for us, turn our lives around, "
        "and testify to Your glorious name. We pray this in the name of the Lord Jesus. Amen!"
    )
    prayer_cn = (
        "親愛的阿爸父，我們感謝讚美祢！感謝祢教導我們：凡帶著需要來到耶穌面前祈求的，就必得著；祂所成就的，甚至超過我們所求所想。"
        "因此我們當效法外邦婦人：帶著信靠的心、謙卑地來到祢面前祈求，把主權交給祢。祢必為我們行奇妙的事，翻轉我們的生命，見證祢榮耀的名。"
        "奉主耶穌的名禱告，阿們！"
    )
    add_textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(2.6), prayer_en, size=14, color=NAVY)
    add_textbox(s, Inches(0.6), Inches(4.2), Inches(12.1), Inches(2.5), prayer_cn, size=14, color=NAVY, font="Microsoft YaHei")

    # ---- 20 Prayer ministry ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY_DARK)
    add_textbox(s, Inches(0.7), Inches(2.3), Inches(12), Inches(0.4),
                "Session 4 · 10 minutes", size=14, color=AMBER)
    add_textbox(s, Inches(0.7), Inches(2.9), Inches(12), Inches(1.0),
                "Prayer Ministry", size=40, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(4.1), Inches(12), Inches(1.2),
                "Two or three people form a prayer group\nand pray for one another’s needs.\n\n兩三人一組，彼此代禱。",
                size=20, color=CREAM, font="Microsoft YaHei")

    total = len(prs.slides)
    # stamp footers on content slides (skip title + section dividers + closing)
    skip_footer = {0, 3, 11, 19}  # 0-indexed
    for i, slide in enumerate(prs.slides):
        if i not in skip_footer:
            footer(slide, i + 1, total)

    out_paths = [
        r"G:\My Drive\Agent\Connect Group\20260919 Mark 7 Connect.pptx",
        r"G:\My Drive\Christianity\ROLCC\20260919 Mark 7 Connect.pptx",
    ]
    for path in out_paths:
        prs.save(path)
        print("Saved:", path)
    print("Total slides:", total)


if __name__ == "__main__":
    build()
